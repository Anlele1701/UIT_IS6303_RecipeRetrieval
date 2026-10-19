"""Build the report .docx (build plan M7) from docs/report/src via the filled Markdown.

    uv run --with pypandoc_binary --with python-docx --with docx2pdf --with pypdf scripts/build_report.py
        [--pdf]    also export a PDF through Word and print the page count (Windows, Word installed)

Steps: fill the numbers (scripts/report_numbers.py), make the Word template (A4, Times New Roman 13, line
spacing 1.3), join the chapters with a title page, convert with pandoc (bundled by pypandoc_binary).
Output: docs/report/A3_bao_cao.docx (and .pdf).
"""

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "docs" / "report"
BUILD = REPORT / "build"
OUT = REPORT / "A3_bao_cao.docx"
TEMPLATE = REPORT / "reference.docx"
sys.path.insert(0, str(ROOT / "scripts"))

TITLE = """---
title: "Trợ lý ẩm thực hội thoại dựa trên RAG, nhận biết ràng buộc của người dùng"
subtitle: "Báo cáo đồ án A3 – IS6303 Thiết kế HTTT thông minh dựa trên truy hồi"
author: "Nhóm 1 – Trường ĐH Công nghệ Thông tin, ĐHQG-HCM"
lang: vi
---

"""
# figure widths in the page (A4 minus margins is about 16 cm); tall diagrams get a narrower width
WIDTH = {"fig_flow": "9cm", "fig_sequence": "15cm", "fig_erd": "16cm", "fig_architecture": "15cm"}


def make_template() -> None:
    import pypandoc
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Cm, Pt, RGBColor
    pypandoc.convert_text("x", "docx", format="md", outputfile=str(TEMPLATE))  # pandoc's own styles to start
    doc = Document(str(TEMPLATE))
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(3), Cm(2)
    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Table", "Image Caption", "Table Caption",
                 "Title", "Subtitle", "Author"):
        if name not in [s.name for s in styles]:
            continue
        st = styles[name]
        st.font.name, st.font.size = "Times New Roman", Pt(13)
        st.font.color.rgb = RGBColor(0, 0, 0)
        pf = st.paragraph_format
        pf.line_spacing_rule, pf.line_spacing = WD_LINE_SPACING.MULTIPLE, 1.3
        pf.space_after = Pt(6)
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13), ("Title", 20), ("Subtitle", 14)):
        if name in [s.name for s in styles]:
            st = styles[name]
            st.font.name, st.font.size, st.font.bold = "Times New Roman", Pt(size), True
            st.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Compact", "Table"):
        if name in [s.name for s in styles]:
            styles[name].font.size = Pt(11)
            styles[name].paragraph_format.line_spacing = 1.0
            styles[name].paragraph_format.space_after = Pt(0)
    for name in ("Source Code", "Verbatim Char"):
        if name in [s.name for s in styles]:
            styles[name].font.name, styles[name].font.size = "Consolas", Pt(8.5)
    if "Image Caption" in [s.name for s in styles]:
        styles["Image Caption"].font.italic = True
        styles["Image Caption"].font.size = Pt(12)
    for st in styles:  # theme fonts (Calibri/Cambria) override the font name: pin every style that is not code
        rpr = st.element.rPr
        if rpr is None or rpr.rFonts is None or st.name in ("Source Code", "Verbatim Char"):
            continue
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rpr.rFonts.attrib.pop(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}", None)
        for attr in ("ascii", "hAnsi", "cs"):
            rpr.rFonts.set(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}", "Times New Roman")
        if st.name.startswith(("Heading", "TOC Heading", "Title", "Subtitle")):
            st.font.color.rgb = RGBColor(0, 0, 0)
    doc.save(str(TEMPLATE))


def finish(path: Path) -> None:
    """Borders on every table (pandoc's table style has none)."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(str(path))
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "808080")
            borders.append(el)
        for old in tbl_pr.findall(qn("w:tblBorders")):
            tbl_pr.remove(old)
        tbl_pr.append(borders)
    doc.save(str(path))


def joined() -> str:
    text = TITLE
    for md in sorted(BUILD.glob("*.md")):
        body = md.read_text(encoding="utf-8")
        # an image gets a width; a figure that does not exist yet stays as a visible note, not a broken link
        def image(m: re.Match) -> str:
            cap, path = m.group(1), m.group(2)
            stem = Path(path).stem
            if not (REPORT / path).exists():
                return f"**[HÌNH CHƯA CÓ: {path}]** *{cap}*"
            return f"![{cap}]({path}){{width={WIDTH.get(stem, '15cm')}}}"
        body = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", image, body)
        text += body + "\n\n"
    return text


def build(pdf: bool) -> None:
    import report_numbers
    import pypandoc
    if report_numbers.main(check=True):
        sys.exit("unfilled numbers: fix docs/report/src first")
    make_template()
    pypandoc.convert_text(joined(), "docx", format="markdown", outputfile=str(OUT),
                          extra_args=[f"--reference-doc={TEMPLATE}", f"--resource-path={REPORT}", "--toc",
                                      "--toc-depth=2", "-M", "toc-title=Mục lục"])
    finish(OUT)
    print(f"-> {OUT.relative_to(ROOT)}")
    if pdf:
        from docx2pdf import convert
        from pypdf import PdfReader
        out_pdf = OUT.with_suffix(".pdf")
        convert(str(OUT), str(out_pdf))
        print(f"-> {out_pdf.relative_to(ROOT)}: {len(PdfReader(str(out_pdf)).pages)} trang")


if __name__ == "__main__":
    build("--pdf" in sys.argv)
